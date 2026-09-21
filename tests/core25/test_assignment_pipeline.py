from pathlib import Path

from openpyxl import load_workbook

from core25 import verify_assignment
from core25.contracts import DecisionState, Domain, Evidence25, Requirement25, Scope


def _requirements():
    return (
        Requirement25(
            requirement_id="SHIFT",
            domain=Domain.ASSIGNMENT,
            text="Продолжительность смены 12 часов",
            scope=Scope.PROJECT_GLOBAL,
            verification_kind="TYPED_VALUE",
            parameter_code="SHIFT_DURATION",
            required_value=12.0,
            unit="час",
            expected_sections=("ТХ",),
        ),
        Requirement25(
            requirement_id="AREA-LAB",
            domain=Domain.ASSIGNMENT,
            text="Площадь здания проборазделки 89,9 м2",
            scope=Scope.OBJECT_SPECIFIC,
            target_object_id="OBJ-LAB",
            verification_kind="TYPED_VALUE",
            parameter_code="AREA",
            required_value=89.9,
            unit="м2",
            expected_sections=("АР",),
        ),
        Requirement25(
            requirement_id="SEMANTIC",
            domain=Domain.ASSIGNMENT,
            text="Обеспечить удобство эксплуатации",
            scope=Scope.UNRESOLVED,
            verification_kind="SEMANTIC_UNSUPPORTED",
        ),
    )


def _evidence():
    return (
        Evidence25(
            evidence_id="E-SHIFT",
            document="ТХ.pdf",
            section="ТХ",
            page=7,
            fragment="Продолжительность смены составляет 12 часов.",
            addressable=True,
            source_kind="PAGE_TEXT",
            metadata={
                "parameter_code": "SHIFT_DURATION",
                "value": 12.0,
                "unit": "час",
            },
        ),
        Evidence25(
            evidence_id="E-DUST-AREA",
            document="АР.pdf",
            section="АР",
            page=12,
            fragment="Модуль обеспыливания. Площадь 23,5 м2.",
            addressable=True,
            source_kind="PAGE_TEXT",
            metadata={
                "owner_name": "Модуль обеспыливания",
                "parameter_code": "AREA",
                "value": 23.5,
                "unit": "м2",
            },
        ),
    )


def _objects():
    return {
        "OBJ-LAB": ("Здание проборазделки", "Проборазделка"),
        "OBJ-DUST": ("Модуль обеспыливания",),
    }


def test_pipeline_rejects_wrong_owner_and_does_not_create_mismatch():
    results = verify_assignment(_requirements(), _evidence(), _objects())
    area = next(x for x in results if x.requirement.requirement_id == "AREA-LAB")

    assert area.decision.state is DecisionState.REVIEW
    assert all(
        binding.owner_id != "OBJ-LAB" or binding.state.value != "BOUND"
        for binding in area.trace.bindings
    )
    assert area.trace.proof.is_categorical is False


def test_pipeline_traces_categorical_result_to_document_page_fragment():
    results = verify_assignment(_requirements(), _evidence(), _objects())
    shift = next(x for x in results if x.requirement.requirement_id == "SHIFT")

    assert shift.decision.state is DecisionState.COMPLIANT
    assert shift.trace.is_categorical_trace_valid() is True
    assert shift.trace.evidence[0].document == "ТХ.pdf"
    assert shift.trace.evidence[0].page == 7
    assert "12 часов" in shift.trace.evidence[0].fragment


def test_unsupported_semantic_requirement_is_review_not_false_success():
    results = verify_assignment(_requirements(), _evidence(), _objects())
    semantic = next(x for x in results if x.requirement.requirement_id == "SEMANTIC")

    assert semantic.decision.state is DecisionState.REVIEW
    assert semantic.trace.proof.is_categorical is False


def _control_assignment_requirements():
    workbook_path = (
        Path(__file__).resolve().parents[2]
        / "validation_reports_150a2"
        / "ExpertCheck_Отчёт_ГИПа_15.0A2.xlsx"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Задание на проектирование"]
    rows = list(sheet.iter_rows(values_only=True))

    header_hints = (
        "id", "требован", "статус", "результ", "доказатель",
        "объект", "показател", "основан",
    )
    scored = []
    for index, row in enumerate(rows[:25]):
        normalized = [str(value or "").strip().casefold() for value in row]
        score = sum(
            1 for value in normalized
            if value and any(hint in value for hint in header_hints)
        )
        scored.append((score, index, normalized))
    score, header_index, normalized_headers = max(scored, default=(0, -1, []))
    assert score >= 2, f"Не найдена строка заголовков контрольного Задания: {scored!r}"

    def _find_header(*needles, exclude=()):
        for index, value in enumerate(normalized_headers):
            if all(needle in value for needle in needles) and not any(
                token in value for token in exclude
            ):
                return index
        return None

    id_index = (
        _find_header("id", "требован")
        if _find_header("id", "требован") is not None
        else _find_header("id")
    )
    text_index = _find_header("требован", exclude=("id", "статус", "результ"))
    if text_index is None:
        text_index = _find_header("строк", "задан")
    assert text_index is not None, (
        "Не найдена колонка текста требования в контрольном отчёте: "
        f"{normalized_headers!r}"
    )

    requirements = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        text = str(row[text_index] or "").strip() if text_index < len(row) else ""
        if not text:
            continue
        requirement_id = ""
        if id_index is not None and id_index < len(row):
            requirement_id = str(row[id_index] or "").strip()
        requirements.append(
            {
                "requirement_id": requirement_id or f"CONTROL-{row_number:03d}",
                "requirement_text": text,
            }
        )
    workbook.close()
    return tuple(requirements), tuple(normalized_headers)


def test_control_package_routes_all_56_assignment_requirements():
    requirements, headers = _control_assignment_requirements()
    assert len(requirements) == 56, (
        f"Ожидалось 56 требований контрольного Задания, получено {len(requirements)}; "
        f"заголовки={headers!r}"
    )

    results = verify_assignment(requirements, (), {})

    assert len(results) == 56
    assert all(result.metadata.get("route_kind") for result in results)
    assert all(result.requirement.scope.value for result in results)
    assert all(
        result.requirement.verification_kind
        or result.metadata.get("route_kind") == "REVIEW_ONLY"
        for result in results
    )
    for result in results:
        if result.decision.state in {DecisionState.COMPLIANT, DecisionState.NONCOMPLIANT}:
            assert result.trace.is_categorical_trace_valid() is True
