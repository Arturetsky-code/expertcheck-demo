from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from core20.report_proof_export import enrich_normative_proof_workbook


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Резюме"
    summary.append(["Показатель", "Значение"])
    summary.append(["Целостность отчёта", "Пройдена"])
    summary.append(["Проверок с независимым AI-консенсусом", 0])

    control = workbook.create_sheet("Контроль отчёта")
    control.append(["Статус", "Проверка", "Результат", "Причина"])
    control.append(["Пройден", "Целостность и согласованность отчёта", "Пройдено", "Ошибок согласованности не выявлено."])

    composition = workbook.create_sheet("Состав проекта")
    composition.append(["Поз.", "Наименование объекта", "Статус", "Основной источник"])
    composition.append(["4.5", "Компрессорная", "Проектируемый", "ПЗУ2, Экспликация/поле генерального плана, стр. 5, nan"])

    normative = workbook.create_sheet("НТД 20.0 — исполнение")
    normative.append(["ID требования", "Требование"])
    normative.append(["R1", "Требование 1"])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _manifest() -> dict:
    return {
        "normative_execution": {
            "semantic_proof_applied": 12,
            "semantic_proof_stale": False,
            "rows": [
                {
                    "requirement_id": "R1",
                    "proof_type": "SEMANTIC_REQUIREMENT",
                    "proof_state": "SEMANTIC_CONSENSUS_PROOF",
                    "retrieval_candidate_count": 4,
                    "semantic_proof": {
                        "judge_verdict": "SUPPORTS",
                        "judge_confidence": 1.0,
                        "judge_provider": "Groq",
                        "judge_model": "openai/gpt-oss-120b",
                        "critic_accept": True,
                        "critic_confidence": 1.0,
                        "critic_provider": "Gemini",
                        "critic_model": "gemini-2.5-flash",
                        "independent": True,
                        "independence_reason": "Провайдеры различаются.",
                        "selected_evidence": [
                            {
                                "evidence_id": "E1",
                                "document": "ПЗ.pdf",
                                "page": 46,
                                "source_locator": "ПЗ.pdf, стр. 46",
                                "fragment": "Содержательный адресный фрагмент проекта.",
                            }
                        ],
                    },
                }
            ],
        }
    }


def _summary_value(workbook, label: str):
    sheet = workbook["Резюме"]
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == label:
            return sheet.cell(row=row, column=2).value
    raise AssertionError(f"summary label not found: {label}")


def test_alpha10_1_export_reconciles_normative_consensus_and_integrity():
    payload = enrich_normative_proof_workbook(_workbook_bytes(), _manifest())
    workbook = load_workbook(BytesIO(payload))

    assert _summary_value(workbook, "Проверок с независимым AI-консенсусом") == 12
    assert _summary_value(workbook, "AI-консенсус — матрица проверки") == 0
    assert _summary_value(workbook, "AI-консенсус — НТД 20.0") == 12

    control = workbook["Контроль отчёта"]
    matching = [
        tuple(control.cell(row=row, column=col).value for col in range(1, 5))
        for row in range(1, control.max_row + 1)
        if control.cell(row=row, column=2).value == "Согласованность AI-консенсуса"
    ]
    assert matching == [
        (
            "Пройден",
            "Согласованность AI-консенсуса",
            "Согласовано",
            "Матрица проверки: 0; НТД 20.0: 12; итог независимого AI-консенсуса: 12.",
        )
    ]


def test_alpha10_1_export_removes_literal_nan_from_composition_source():
    payload = enrich_normative_proof_workbook(_workbook_bytes(), _manifest())
    workbook = load_workbook(BytesIO(payload))
    source = workbook["Состав проекта"]["D2"].value

    assert source == "ПЗУ2, Экспликация/поле генерального плана, стр. 5"
    assert "nan" not in source.casefold()


def test_alpha10_1_export_is_idempotent_for_consensus_and_proof_columns():
    first = enrich_normative_proof_workbook(_workbook_bytes(), _manifest())
    second = enrich_normative_proof_workbook(first, _manifest())
    workbook = load_workbook(BytesIO(second))

    assert _summary_value(workbook, "Проверок с независимым AI-консенсусом") == 12
    summary = workbook["Резюме"]
    assert sum(summary.cell(row=row, column=1).value == "AI-консенсус — матрица проверки" for row in range(1, summary.max_row + 1)) == 1
    assert sum(summary.cell(row=row, column=1).value == "AI-консенсус — НТД 20.0" for row in range(1, summary.max_row + 1)) == 1

    normative = workbook["НТД 20.0 — исполнение"]
    headers = [cell.value for cell in normative[1]]
    assert headers.count("Выбранные доказательства") == 1
    assert headers.count("Примечание к уверенности") == 1
