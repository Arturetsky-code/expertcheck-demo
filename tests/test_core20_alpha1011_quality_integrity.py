from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from core.quality_integrity_1011_patch import (
    _sanitize_assembly_row,
    clean_source_text,
    compact_source_clean,
)
from core20.quality_integrity_1011 import (
    candidate_payloads_without_toc,
    is_toc_like_text,
    reconcile_report_consensus,
)


def test_alpha1011_object_source_removes_literal_nan():
    source = "ПЗУ2, Экспликация/поле генерального плана, стр. 5, nan"
    assert clean_source_text(source) == "ПЗУ2, Экспликация/поле генерального плана, стр. 5"
    assert compact_source_clean({
        "document_type": "ПЗУ2",
        "section": "Экспликация/поле генерального плана",
        "page": 5,
        "table": float("nan"),
    }) == "ПЗУ2, Экспликация/поле генерального плана, стр. 5"

    cleaned = _sanitize_assembly_row({
        "Канонический источник": source,
        "_evidence": [{
            "document_type": "ПЗУ2",
            "section": "Экспликация/поле генерального плана",
            "page": 5,
            "table": "nan",
            "row": float("nan"),
        }],
    })
    assert cleaned["Канонический источник"] == "ПЗУ2, Экспликация/поле генерального плана, стр. 5"
    assert cleaned["_evidence"][0]["table"] == ""
    assert cleaned["_evidence"][0]["row"] == ""


def test_alpha1011_toc_page_is_not_normative_proof_candidate():
    toc = (
        "Лист 3 Содержание 1 Характеристика земельного участка ............ 6 "
        "2 Обоснование санитарно-защитных зон ............ 12 "
        "5 Обоснование и описание решений по инженерной подготовке территории ............ 19 "
        "9 Обоснование схем транспортных коммуникаций ............ 28"
    )
    substantive = (
        "5 Обоснование и описание решений по инженерной подготовке территории. "
        "Проектом предусмотрен организованный отвод поверхностных вод и инженерная подготовка площадки."
    )
    assert is_toc_like_text(toc) is True
    assert is_toc_like_text(substantive) is False

    ranked = [
        (7, 1.0, len(toc), {"document": "ПЗУ1.pdf", "page": 4, "text": toc, "document_type": "ПЗУ"}, ["обоснование", "решений"]),
        (5, 0.7, len(substantive), {"document": "ПЗУ1.pdf", "page": 20, "text": substantive, "document_type": "ПЗУ"}, ["обоснование", "решений"]),
    ]
    result = candidate_payloads_without_toc(ranked, "PP87-12-E", limit=4)
    assert len(result) == 1
    assert result[0]["page"] == 20
    assert result[0]["evidence_quality"] == "SUBSTANTIVE_PAGE"
    assert result[0]["semantic_eligible"] is True


def _workbook_payload() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Резюме"
    summary.append(["Показатель", "Значение"])
    summary.append(["Проверок с независимым AI-консенсусом", 0])

    control = workbook.create_sheet("Контроль отчёта")
    control.append(["Статус", "Проверка", "Результат", "Причина"])
    control.append([
        "Пройден",
        "Целостность и согласованность отчёта",
        "Пройдено",
        "Ошибок согласованности не выявлено.",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _read_consensus(payload: bytes) -> tuple[int, str]:
    workbook = load_workbook(BytesIO(payload), data_only=False)
    summary = workbook["Резюме"]
    control = workbook["Контроль отчёта"]
    return int(summary.cell(row=2, column=2).value or 0), str(control.cell(row=2, column=4).value or "")


def test_alpha1011_export_consensus_includes_normative_and_is_idempotent():
    manifest = {
        "normative_execution": {
            "semantic_proof_applied": 12,
            "semantic_proof_stale": False,
        }
    }
    first = reconcile_report_consensus(_workbook_payload(), manifest)
    total, reason = _read_consensus(first)
    assert total == 12
    assert "матрица 0, НТД 12, итого 12" in reason
    assert "Alpha 10.1.1: AI-консенсус сверен" in reason

    second = reconcile_report_consensus(first, manifest)
    total_again, reason_again = _read_consensus(second)
    assert total_again == 12
    assert reason_again == reason


def test_alpha1011_stale_normative_proof_is_not_added_to_consensus():
    manifest = {
        "normative_execution": {
            "semantic_proof_applied": 12,
            "semantic_proof_stale": True,
        }
    }
    payload = reconcile_report_consensus(_workbook_payload(), manifest)
    total, reason = _read_consensus(payload)
    assert total == 0
    assert reason == "Ошибок согласованности не выявлено."
