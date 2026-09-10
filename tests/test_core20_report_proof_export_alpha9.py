from io import BytesIO

from openpyxl import Workbook, load_workbook

from core20.report_proof_export import (
    enrich_normative_proof_workbook,
    reconcile_project_data_contract_workbook,
)


def _workbook_bytes():
    wb=Workbook()
    ws=wb.active
    ws.title="НТД 20.0 — исполнение"
    ws.append(["ID требования","НТД","Результат"])
    ws.append(["PP87-X-SEM","ПП РФ №87","Подтверждено"])
    out=BytesIO(); wb.save(out); return out.getvalue()


def _manifest():
    return {
        "normative_execution":{
            "rows":[{
                "requirement_id":"PP87-X-SEM",
                "proof_type":"SEMANTIC_REQUIREMENT",
                "proof_state":"SEMANTIC_CONSENSUS_PROOF",
                "retrieval_candidate_count":2,
                "semantic_proof":{
                    "judge_verdict":"SUPPORTS",
                    "judge_confidence":0.96,
                    "judge_provider":"Groq",
                    "judge_model":"model-a",
                    "critic_accept":True,
                    "critic_confidence":0.91,
                    "critic_provider":"Gemini",
                    "critic_model":"model-b",
                    "independent":True,
                    "independence_reason":"Провайдеры и модели различаются.",
                    "selected_evidence":[
                        {"evidence_id":"NORM-E-1","document":"ПЗУ.pdf","page":10,"source_locator":"ПЗУ.pdf, стр. 10"},
                        {"evidence_id":"NORM-E-2","document":"ПЗУ.pdf","page":11,"source_locator":"ПЗУ.pdf, стр. 11"},
                    ],
                },
            }]
        }
    }


def test_alpha9_export_appends_proof_and_judge_critic_trace():
    enriched=enrich_normative_proof_workbook(_workbook_bytes(),_manifest())
    wb=load_workbook(BytesIO(enriched),data_only=True)
    ws=wb["НТД 20.0 — исполнение"]
    headers={cell.value:cell.column for cell in ws[1]}
    assert "Тип доказательства" in headers
    assert "Состояние доказательства" in headers
    assert "Решение проверяющей модели" in headers
    assert "Контрольная модель приняла" in headers
    assert "Выбранные доказательства" in headers
    assert ws.cell(2,headers["Тип доказательства"]).value=="Смысловое выполнение требования"
    assert ws.cell(2,headers["Смысловое доказательство применено"]).value=="Да"
    assert ws.cell(2,headers["Решение проверяющей модели"]).value=="Подтверждает"
    assert ws.cell(2,headers["Контрольная модель приняла"]).value=="Да"
    trace=ws.cell(2,headers["Выбранные доказательства"]).value
    assert "ПЗУ.pdf, стр. 10" in trace
    assert "ПЗУ.pdf, стр. 11" in trace


def test_alpha9_export_is_noop_when_execution_sheet_absent():
    wb=Workbook(); wb.active.title="Резюме"
    out=BytesIO(); wb.save(out)
    original=out.getvalue()
    enriched=enrich_normative_proof_workbook(original,_manifest())
    assert enriched==original


def test_alpha9_report_contract_summary_uses_project_checkpoint_not_export_repairs():
    wb=Workbook()
    summary=wb.active; summary.title="Резюме"
    summary.append(["Показатель","Значение"])
    summary.append(["Контракт данных 18.0","REPAIRED"])
    summary.append(["Исправлено значений контрактом",146346])
    summary.append(["Отпечаток результата","export-fingerprint"])
    control=wb.create_sheet("Контроль данных")
    control.append(["Показатель","Значение"])
    control.append(["Версия контракта","18.0-project-data-contract-v1"])
    control.append(["Статус","REPAIRED"])
    control.append(["Исправлено значений",146346])
    control.append(["Документов",12])
    control.append(["Находок",1409])
    control.append(["Сверок",99])
    control.append(["Отпечаток результата","export-fingerprint"])
    control.append(["Исправление: non_finite_number",146346])
    out=BytesIO(); wb.save(out)

    contract={
        "version":"18.0-project-data-contract-v1",
        "status":"REPAIRED",
        "repairs":680,
        "counts":{"documents":12,"findings":1409,"comparisons":99},
        "result_identity_fingerprint":"project-fingerprint",
    }
    reconciled=reconcile_project_data_contract_workbook(out.getvalue(),contract)
    wb=load_workbook(BytesIO(reconciled),data_only=True)
    summary=wb["Резюме"]
    values={summary.cell(row=i,column=1).value:summary.cell(row=i,column=2).value for i in range(2,summary.max_row+1)}
    assert values["Исправлено значений контрактом"]==680
    assert values["Отпечаток результата"]=="project-fingerprint"

    control=wb["Контроль данных"]
    values={control.cell(row=i,column=1).value:control.cell(row=i,column=2).value for i in range(2,control.max_row+1)}
    assert values["Исправлено значений"]==680
    assert values["Экспортный контроль: исправлено значений"]==146346
    assert values["Экспортный контроль: отпечаток результата"]=="export-fingerprint"
    assert "Экспортное исправление: non_finite_number" in values
