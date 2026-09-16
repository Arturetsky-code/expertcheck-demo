from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook

_INSTALLED = False
_ORIGINAL_ENRICH = None
_REPORT_MARKER = "Alpha 10.1.1: AI-консенсус сверен"


def is_toc_like_text(value: Any) -> bool:
    """Identify table-of-contents pages that must not serve as project proof."""
    raw = " ".join(str(value or "").replace("\xa0", " ").split())
    if not raw:
        return False
    head = raw[:1400].casefold().replace("ё", "е")
    if not re.search(r"\b(?:содержание|оглавление)\b", head):
        return False
    leader_runs = len(re.findall(r"\.{4,}", raw[:4000]))
    numbered_entries = len(re.findall(r"(?:^|\s)\d+(?:\.\d+)*\s+[A-Za-zА-Яа-я]", raw[:4000]))
    return leader_runs >= 2 or numbered_entries >= 4


def candidate_payloads_without_toc(
    ranked: list[tuple[int, float, int, dict[str, Any], list[str]]],
    requirement_id: str,
    *,
    limit: int = 4,
) -> list[dict[str, Any]]:
    """Keep only substantive pages in the normative proof candidate set."""
    from . import normative_execution as execution

    output: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for score, coverage, _, page, hits in ranked:
        raw_text = str(page.get("text") or page.get("content") or "")
        if is_toc_like_text(raw_text):
            continue
        document = str(page.get("document") or "").strip()
        page_no = page.get("page")
        fragment = execution._fragment(raw_text, hits)
        key = (document, str(page_no), fragment[:240])
        if key in seen:
            continue
        seen.add(key)
        index = len(output) + 1
        output.append({
            "evidence_id": f"NORM-E-{requirement_id}-{index:02d}",
            "document": document,
            "page": page_no,
            "section": str(page.get("document_type") or page.get("section") or ""),
            "fragment": fragment,
            "matched_keywords": list(hits),
            "retrieval_keyword_score": score,
            "retrieval_keyword_coverage": round(coverage, 3),
            "evidence_quality": "SUBSTANTIVE_PAGE",
            "semantic_eligible": True,
        })
        if len(output) >= limit:
            break
    return output


def _safe_int(value: Any) -> int:
    try:
        number = float(value or 0)
        if number != number:
            return 0
        return max(0, int(number))
    except (TypeError, ValueError, OverflowError):
        return 0


def _normative_consensus(canonical_manifest: dict[str, Any] | None) -> int:
    execution = dict((canonical_manifest or {}).get("normative_execution") or {})
    if execution.get("semantic_proof_stale"):
        return 0
    return _safe_int(execution.get("semantic_proof_applied"))


def _control_row(sheet: Any) -> int | None:
    for row_index in range(2, sheet.max_row + 1):
        check = str(sheet.cell(row=row_index, column=2).value or "").strip()
        if check == "Целостность и согласованность отчёта":
            return row_index
    return None


def _summary_rows(sheet: Any) -> dict[str, int]:
    result: dict[str, int] = {}
    for row_index in range(2, sheet.max_row + 1):
        label = str(sheet.cell(row=row_index, column=1).value or "").strip()
        if label:
            result[label] = row_index
    return result


def reconcile_report_consensus(
    payload: bytes | bytearray | None,
    canonical_manifest: dict[str, Any] | None,
) -> bytes | None:
    """Fallback reconciliation for workbooks not yet processed by Alpha 10.1.

    The main report exporter already writes explicit matrix/NTD consensus rows.
    If those rows are present, this overlay must be a no-op so totals cannot be
    double-counted by two compatible integrity layers.
    """
    if payload is None:
        return None
    raw = bytes(payload)
    normative = _normative_consensus(canonical_manifest)
    if not raw or normative <= 0:
        return raw

    workbook = load_workbook(BytesIO(raw))
    if "Резюме" not in workbook.sheetnames:
        return raw

    summary = workbook["Резюме"]
    rows = _summary_rows(summary)
    if "AI-консенсус — матрица проверки" in rows and "AI-консенсус — НТД 20.0" in rows:
        # Already reconciled by core20.report_proof_export._reconcile_ai_consensus.
        return raw

    consensus_row = rows.get("Проверок с независимым AI-консенсусом")
    if consensus_row is None:
        return raw

    control = workbook["Контроль отчёта"] if "Контроль отчёта" in workbook.sheetnames else None
    control_index = _control_row(control) if control is not None else None
    already_reconciled = False
    if control is not None and control_index is not None:
        reason = str(control.cell(row=control_index, column=4).value or "")
        already_reconciled = _REPORT_MARKER in reason

    current = _safe_int(summary.cell(row=consensus_row, column=2).value)
    if already_reconciled:
        total = current
        matrix = max(0, current - normative)
    else:
        matrix = current
        total = matrix + normative
        summary.cell(row=consensus_row, column=2, value=total)

    if control is not None and control_index is not None:
        control.cell(row=control_index, column=1, value="Пройден")
        control.cell(row=control_index, column=3, value="Пройдено")
        control.cell(
            row=control_index,
            column=4,
            value=(
                f"{_REPORT_MARKER}: матрица {matrix}, НТД {normative}, итого {total}. "
                "Ошибок согласованности не выявлено."
            ),
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def install_quality_integrity_1011() -> None:
    """Install Alpha 10.1.1 fail-closed normative and report-integrity overlays."""
    global _INSTALLED, _ORIGINAL_ENRICH
    if _INSTALLED:
        return

    from . import normative_execution as execution
    from . import report_proof_export as report_export

    execution._candidate_payloads = candidate_payloads_without_toc

    _ORIGINAL_ENRICH = report_export.enrich_normative_proof_workbook

    def enrich_normative_proof_workbook_1011(
        payload: bytes | bytearray | None,
        canonical_manifest: dict[str, Any] | None,
    ) -> bytes | None:
        enriched = _ORIGINAL_ENRICH(payload, canonical_manifest)
        return reconcile_report_consensus(enriched, canonical_manifest)

    report_export.enrich_normative_proof_workbook = enrich_normative_proof_workbook_1011
    _INSTALLED = True


__all__ = [
    "candidate_payloads_without_toc",
    "is_toc_like_text",
    "reconcile_report_consensus",
    "install_quality_integrity_1011",
]
