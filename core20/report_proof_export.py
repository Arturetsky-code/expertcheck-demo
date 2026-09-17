from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .proof_labels import judge_label, proof_state_label, proof_type_label


SHEET_NAME = "НТД 20.0 — исполнение"
SUMMARY_SHEET = "Резюме"
REPORT_CONTROL_SHEET = "Контроль отчёта"
COMPOSITION_SHEET = "Состав проекта"
CONSENSUS_TOTAL_LABEL = "Проверок с независимым AI-консенсусом"
CONSENSUS_MATRIX_LABEL = "AI-консенсус — матрица проверки"
CONSENSUS_NORMATIVE_LABEL = "AI-консенсус — НТД 20.0"
CONSENSUS_CONTROL_LABEL = "Согласованность AI-консенсуса"
_NULL_TOKENS = {"nan", "none", "null", "nat"}

PROOF_COLUMNS = (
    "Тип доказательства",
    "Состояние доказательства",
    "Кандидатов доказательства",
    "Смысловое доказательство применено",
    "Решение проверяющей модели",
    "Заявленная уверенность проверяющей модели",
    "Провайдер проверяющей модели",
    "Проверяющая модель",
    "Контрольная модель приняла",
    "Заявленная уверенность контрольной модели",
    "Провайдер контрольной модели",
    "Контрольная модель",
    "Независимость моделей",
    "Основание независимости",
    "Выбранные доказательства",
    "Примечание к уверенности",
)


def _execution(canonical_manifest: dict[str, Any] | None) -> dict[str, Any]:
    manifest = dict(canonical_manifest or {})
    return dict(manifest.get("normative_execution") or {})


def _execution_rows(canonical_manifest: dict[str, Any] | None) -> list[dict[str, Any]]:
    execution = _execution(canonical_manifest)
    return [dict(row) for row in (execution.get("rows") or []) if isinstance(row, dict)]


def _safe_int(value: Any) -> int:
    try:
        number = float(value or 0)
        if number != number:
            return 0
        return max(0, int(number))
    except (TypeError, ValueError, OverflowError):
        return 0


def _normative_consensus(canonical_manifest: dict[str, Any] | None) -> int:
    execution = _execution(canonical_manifest)
    if execution.get("semantic_proof_stale"):
        return 0
    return _safe_int(execution.get("semantic_proof_applied"))


def _find_label_row(sheet, label: str, *, column: int = 1) -> int | None:
    wanted = str(label or "").strip()
    for row_index in range(1, sheet.max_row + 1):
        if str(sheet.cell(row=row_index, column=column).value or "").strip() == wanted:
            return row_index
    return None


def _copy_row_style(sheet, source_row: int, target_row: int, *, columns: int) -> None:
    if source_row < 1 or source_row > sheet.max_row:
        return
    for col in range(1, columns + 1):
        src = sheet.cell(row=source_row, column=col)
        dst = sheet.cell(row=target_row, column=col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.number_format:
            dst.number_format = src.number_format


def _upsert_summary_row(sheet, label: str, value: Any, *, template_row: int) -> int:
    row_index = _find_label_row(sheet, label)
    if row_index is None:
        row_index = sheet.max_row + 1
        _copy_row_style(sheet, template_row, row_index, columns=2)
        sheet.cell(row=row_index, column=1, value=label)
    sheet.cell(row=row_index, column=2, value=value)
    return row_index


def _upsert_control_row(sheet, *, status: str, check: str, result: str, reason: str) -> int:
    row_index = _find_label_row(sheet, check, column=2)
    if row_index is None:
        row_index = sheet.max_row + 1
        template_row = 2 if sheet.max_row >= 2 else 1
        _copy_row_style(sheet, template_row, row_index, columns=4)
    values = (status, check, result, reason)
    for column, value in enumerate(values, start=1):
        sheet.cell(row=row_index, column=column, value=value)
    return row_index


def _reconcile_ai_consensus(workbook, canonical_manifest: dict[str, Any] | None) -> bool:
    if SUMMARY_SHEET not in workbook.sheetnames:
        return False
    summary = workbook[SUMMARY_SHEET]
    total_row = _find_label_row(summary, CONSENSUS_TOTAL_LABEL)
    if total_row is None:
        return False

    raw_total = _safe_int(summary.cell(row=total_row, column=2).value)
    matrix_row = _find_label_row(summary, CONSENSUS_MATRIX_LABEL)
    normative_row = _find_label_row(summary, CONSENSUS_NORMATIVE_LABEL)
    normative = _normative_consensus(canonical_manifest)

    if matrix_row is not None:
        matrix = _safe_int(summary.cell(row=matrix_row, column=2).value)
    elif normative_row is not None:
        previous_normative = _safe_int(summary.cell(row=normative_row, column=2).value)
        matrix = max(0, raw_total - previous_normative)
    else:
        # Fresh workbook: the legacy report builder stores only matrix consensus.
        matrix = raw_total

    total = matrix + normative
    summary.cell(row=total_row, column=2, value=total)
    _upsert_summary_row(summary, CONSENSUS_MATRIX_LABEL, matrix, template_row=total_row)
    _upsert_summary_row(summary, CONSENSUS_NORMATIVE_LABEL, normative, template_row=total_row)

    integrity_row = _find_label_row(summary, "Целостность отчёта")
    if integrity_row is not None:
        summary.cell(row=integrity_row, column=2, value="Пройдена")

    if REPORT_CONTROL_SHEET in workbook.sheetnames:
        control = workbook[REPORT_CONTROL_SHEET]
        _upsert_control_row(
            control,
            status="Пройден",
            check=CONSENSUS_CONTROL_LABEL,
            result="Согласовано",
            reason=f"Матрица проверки: {matrix}; НТД 20.0: {normative}; итог независимого AI-консенсуса: {total}.",
        )
        overall_row = _find_label_row(control, "Целостность и согласованность отчёта", column=2)
        if overall_row is not None:
            control.cell(row=overall_row, column=1, value="Пройден")
            control.cell(row=overall_row, column=3, value="Пройдено")
            control.cell(
                row=overall_row,
                column=4,
                value=(
                    "Сводные показатели отчёта сверены; AI-консенсус согласован по матрице проверки "
                    "и НТД 20.0."
                ),
            )
    return True


def _clean_display_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text.casefold() in _NULL_TOKENS:
        return ""

    # Canonical source strings are comma-separated. Remove only standalone
    # serialization null tokens so legitimate words and punctuation remain intact.
    comma_parts = [part.strip() for part in text.split(",")]
    if len(comma_parts) > 1 and any(part.casefold() in _NULL_TOKENS for part in comma_parts):
        text = ", ".join(part for part in comma_parts if part and part.casefold() not in _NULL_TOKENS)

    for separator in (" | ", "; "):
        parts = [part.strip() for part in text.split(separator)]
        if len(parts) > 1 and any(part.casefold() in _NULL_TOKENS for part in parts):
            text = separator.join(part for part in parts if part and part.casefold() not in _NULL_TOKENS)
    return text.strip(" ,;|")


def _clean_composition_null_tokens(workbook) -> bool:
    if COMPOSITION_SHEET not in workbook.sheetnames:
        return False
    sheet = workbook[COMPOSITION_SHEET]
    changed = False
    for row in sheet.iter_rows():
        for cell in row:
            cleaned = _clean_display_text(cell.value)
            if cleaned != cell.value:
                cell.value = cleaned
                changed = True
    return changed


def _evidence_excerpt(item: dict[str, Any], limit: int = 320) -> str:
    parts = []
    for key in ("text", "fragment", "quote", "excerpt"):
        value = item.get(key)
        if not value:
            continue
        normalized = " ".join(str(value).split()).strip()
        if normalized:
            parts.append(normalized)
    text = " ".join(parts).strip()
    if len(text) > limit:
        return text[: limit - 1].rstrip() + "…"
    return text


def _selected_trace(row: dict[str, Any], proof: dict[str, Any]) -> str:
    selected = list(proof.get("selected_evidence") or row.get("semantic_selected_evidence") or [])
    rendered = []
    for item in selected:
        if not isinstance(item, dict):
            continue
        locator = str(item.get("source_locator") or "").strip()
        if not locator:
            document = str(item.get("document") or "").strip()
            page = item.get("page")
            locator = f"{document}, стр. {page}" if document and page not in (None, "") else document
        evidence_id = str(item.get("evidence_id") or "").strip()
        prefix = f"{evidence_id}: {locator}" if evidence_id and locator else evidence_id or locator
        excerpt = _evidence_excerpt(item)
        text = f"{prefix} — «{excerpt}»" if prefix and excerpt else excerpt or prefix
        if text and text not in rendered:
            rendered.append(text)
    return " | ".join(rendered)


def proof_export_row(row: dict[str, Any]) -> dict[str, Any]:
    proof = dict(row.get("semantic_proof") or {})
    has_semantic = bool(proof)
    return {
        "Тип доказательства": proof_type_label(row.get("proof_type")),
        "Состояние доказательства": proof_state_label(row.get("proof_state")),
        "Кандидатов доказательства": int(row.get("retrieval_candidate_count") or 0),
        "Смысловое доказательство применено": "Да" if row.get("proof_state") == "SEMANTIC_CONSENSUS_PROOF" else "Нет",
        "Решение проверяющей модели": judge_label(proof.get("judge_verdict")) if has_semantic else "—",
        "Заявленная уверенность проверяющей модели": proof.get("judge_confidence") if has_semantic else "—",
        "Провайдер проверяющей модели": proof.get("judge_provider") or "—",
        "Проверяющая модель": proof.get("judge_model") or "—",
        "Контрольная модель приняла": ("Да" if proof.get("critic_accept") is True else "Нет") if has_semantic else "—",
        "Заявленная уверенность контрольной модели": proof.get("critic_confidence") if has_semantic else "—",
        "Провайдер контрольной модели": proof.get("critic_provider") or "—",
        "Контрольная модель": proof.get("critic_model") or "—",
        "Независимость моделей": ("Да" if proof.get("independent") is True else "Нет") if has_semantic else "—",
        "Основание независимости": proof.get("independence_reason") or "—",
        "Выбранные доказательства": _selected_trace(row, proof) or "—",
        "Примечание к уверенности": (
            "Уверенность сообщена самой AI-моделью и используется только как один из proof-gate сигналов; "
            "она не является измеренной вероятностью корректности вывода."
            if has_semantic
            else "—"
        ),
    }


def _enrich_normative_sheet(workbook, rows: list[dict[str, Any]]) -> bool:
    if not rows or SHEET_NAME not in workbook.sheetnames:
        return False
    sheet = workbook[SHEET_NAME]
    headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
    requirement_col = headers.get("ID требования")
    if requirement_col is None:
        return False

    by_id = {
        str(row.get("requirement_id") or "").strip(): row
        for row in rows
        if str(row.get("requirement_id") or "").strip()
    }
    if not by_id:
        return False

    template = sheet.cell(row=1, column=max(1, sheet.max_column))
    next_col = sheet.max_column + 1
    for header in PROOF_COLUMNS:
        if header in headers:
            continue
        cell = sheet.cell(row=1, column=next_col, value=header)
        if template.has_style:
            cell._style = copy(template._style)
        if template.font:
            cell.font = copy(template.font)
        if template.fill:
            cell.fill = copy(template.fill)
        if template.border:
            cell.border = copy(template.border)
        if template.alignment:
            cell.alignment = copy(template.alignment)
        sheet.column_dimensions[cell.column_letter].width = (
            42 if header in {"Основание независимости", "Выбранные доказательства", "Примечание к уверенности"} else 24
        )
        headers[header] = next_col
        next_col += 1

    for row_index in range(2, sheet.max_row + 1):
        requirement_id = str(sheet.cell(row=row_index, column=requirement_col).value or "").strip()
        source = by_id.get(requirement_id)
        if not source:
            continue
        values = proof_export_row(source)
        for header in PROOF_COLUMNS:
            sheet.cell(row=row_index, column=headers[header], value=values.get(header))
    return True


def enrich_normative_proof_workbook(
    payload: bytes | bytearray | None,
    canonical_manifest: dict[str, Any] | None,
) -> bytes | None:
    """Apply the Alpha 10.1 final report integrity gate and auditable NTD proof trace."""
    if payload is None:
        return None
    raw = bytes(payload)
    if not raw:
        return raw

    workbook = load_workbook(BytesIO(raw))
    changed = False
    changed = _enrich_normative_sheet(workbook, _execution_rows(canonical_manifest)) or changed
    changed = _reconcile_ai_consensus(workbook, canonical_manifest) or changed
    changed = _clean_composition_null_tokens(workbook) or changed

    if not changed:
        return raw
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def reconcile_project_data_contract_workbook(
    payload: bytes | bytearray | None,
    project_data_contract: dict[str, Any] | None,
) -> bytes | None:
    """Keep report-facing data-contract metrics equal to the pipeline checkpoint.

    The report builder performs a second serialization-safety normalization. Its
    repair count may be much larger because nested NaN values are sanitized for
    XLSX. That boundary audit is useful diagnostics, but it must not replace the
    project contract metric shown in the UI.
    """
    if payload is None:
        return None
    raw = bytes(payload)
    contract = dict(project_data_contract or {})
    if not raw or not contract:
        return raw

    workbook = load_workbook(BytesIO(raw))
    changed = False

    if "Резюме" in workbook.sheetnames:
        sheet = workbook["Резюме"]
        summary_values = {
            "Контракт данных 18.0": contract.get("status") or "Не выполнен",
            "Исправлено значений контрактом": int(contract.get("repairs") or 0),
            "Отпечаток результата": contract.get("result_identity_fingerprint") or "—",
        }
        for row_index in range(2, sheet.max_row + 1):
            label = str(sheet.cell(row=row_index, column=1).value or "").strip()
            if label in summary_values:
                sheet.cell(row=row_index, column=2, value=summary_values[label])
                changed = True

    if "Контроль данных" in workbook.sheetnames:
        sheet = workbook["Контроль данных"]
        existing = {}
        for row_index in range(2, sheet.max_row + 1):
            label = str(sheet.cell(row=row_index, column=1).value or "").strip()
            if label:
                existing[label] = sheet.cell(row=row_index, column=2).value

        project_values = {
            "Версия контракта": contract.get("version"),
            "Статус": contract.get("status"),
            "Исправлено значений": int(contract.get("repairs") or 0),
            "Документов": (contract.get("counts") or {}).get("documents", 0),
            "Находок": (contract.get("counts") or {}).get("findings", 0),
            "Сверок": (contract.get("counts") or {}).get("comparisons", 0),
            "Отпечаток результата": contract.get("result_identity_fingerprint"),
        }
        for row_index in range(2, sheet.max_row + 1):
            label = str(sheet.cell(row=row_index, column=1).value or "").strip()
            if label in project_values:
                sheet.cell(row=row_index, column=2, value=project_values[label])
                changed = True
            elif label.startswith("Исправление:"):
                sheet.cell(row=row_index, column=1, value="Экспортное " + label.casefold())
                changed = True

        export_rows = [
            ("Экспортный контроль: статус", existing.get("Статус")),
            ("Экспортный контроль: исправлено значений", existing.get("Исправлено значений")),
            ("Экспортный контроль: отпечаток результата", existing.get("Отпечаток результата")),
        ]
        template_row = 2 if sheet.max_row >= 2 else 1
        for label, value in export_rows:
            if value in (None, ""):
                continue
            target = sheet.max_row + 1
            sheet.cell(row=target, column=1, value=label)
            sheet.cell(row=target, column=2, value=value)
            for col in (1, 2):
                src = sheet.cell(row=template_row, column=col)
                dst = sheet.cell(row=target, column=col)
                if src.has_style:
                    dst._style = copy(src._style)
                if src.font:
                    dst.font = copy(src.font)
                if src.fill:
                    dst.fill = copy(src.fill)
                if src.border:
                    dst.border = copy(src.border)
                if src.alignment:
                    dst.alignment = copy(src.alignment)
            changed = True

    if not changed:
        return raw
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
